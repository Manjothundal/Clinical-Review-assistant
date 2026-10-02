# -*- coding: utf-8 -*-
"""
Writes the findings sheet.

The output is deliberately the same shape as the sheet the team already fills in
by hand, so a reviewer can work in it without learning anything new:

    tab "Findings"  - the five columns, navy header, filter row, frozen panes
    tab "Read me"   - what was reviewed, what was checked, what was not

The second tab matters. A reviewer has to be able to see what the run did NOT
look at, or they will assume it looked at everything.
"""
import os
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .findings import COLUMNS, INTERNAL, to_dataframe

NAVY = "17336E"
STRIPE = "EEF2F8"
GREY = "C2CCDA"

WIDTHS = [14, 40, 34, 62, 78]
INTERNAL_WIDTHS = [14, 11, 26, 12, 14, 11]


def _style_sheet(ws, n_cols, n_rows, widths):
    thin = Side(style="thin", color=GREY)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="left")
        cell.border = border
    ws.row_dimensions[1].height = 30

    for row in range(2, n_rows + 2):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="left")
            cell.border = border
            cell.font = Font(size=10)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=STRIPE)
        ws.cell(row=row, column=1).font = Font(size=10, bold=True)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "C2"
    if n_rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(n_cols), n_rows + 1)


def write_findings(findings, path, study="", checks_run=None, not_checked=None,
                   internal=False, candidates=None):
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)

    df = to_dataframe(findings, internal=internal)
    cols = COLUMNS + INTERNAL if internal else COLUMNS
    widths = WIDTHS + INTERNAL_WIDTHS if internal else WIDTHS

    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append(cols)
    for _, r in df.iterrows():
        ws.append([r[c] for c in cols])
    _style_sheet(ws, len(cols), len(df), widths)

    # ---- things the code noticed but could not prove ---------------------
    if candidates:
        ws3 = wb.create_sheet("To look at")
        ws3.append(cols)
        for _, r in to_dataframe(candidates, internal=internal).iterrows():
            ws3.append([r[c] for c in cols])
        _style_sheet(ws3, len(cols), len(candidates), widths)

    # ---- the Read me tab -------------------------------------------------
    ws2 = wb.create_sheet("Read me")
    by_area = {}
    for f in findings:
        by_area[f.area] = by_area.get(f.area, 0) + 1
    by_sev = {}
    for f in findings:
        by_sev[f.severity] = by_sev.get(f.severity, 0) + 1

    notes = [
        ["Study", study],
        ["Run on", datetime.datetime.now().strftime("%d %b %Y %H:%M")],
        ["Findings", str(len(findings))],
        ["To look at", str(len(candidates or []))],
        ["", ""],
        ["How to read this", ""],
        ["Evidence",
         "Every number in the Evidence column was produced by code reading the "
         "delivered outputs. No number here was written by a model."],
        ["What is wrong",
         "The wording. Where a model was used it was allowed to choose the "
         "words only, never the numbers, and any number it wrote that is not in "
         "the Evidence column causes the finding to be dropped."],
        ["Confidence",
         "A finding marked candidate was found by code that cannot be sure on "
         "its own, and was kept or dropped by judgement. Everything else was "
         "proved by code."],
        ["", ""],
        ["By severity", ""],
    ]
    for s in ("Critical", "Major", "Minor"):
        if s in by_sev:
            notes.append([s, str(by_sev[s])])
    if candidates:
        notes += [["", ""],
                  ["To look at",
                   "%d item(s) are on their own tab. Code noticed them but "
                   "cannot prove them on its own, and no judgement step ran. "
                   "They are not findings yet." % len(candidates)]]
    notes += [["", ""], ["By area", ""]]
    for a, n in sorted(by_area.items(), key=lambda kv: -kv[1]):
        notes.append([a, str(n)])

    if checks_run:
        notes += [["", ""], ["Checks that ran", ""]]
        for c in sorted(set(checks_run)):
            notes.append([c, ""])
    if not_checked:
        notes += [["", ""], ["NOT checked in this run", ""]]
        for c in not_checked:
            notes.append(["", c])

    for row in notes:
        ws2.append(row)
    for row in range(1, ws2.max_row + 1):
        a, b = ws2.cell(row=row, column=1), ws2.cell(row=row, column=2)
        a.font = Font(bold=True, size=10)
        a.alignment = Alignment(vertical="top")
        b.font = Font(size=10)
        b.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 95

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# the full scan: one tab per layer
# ---------------------------------------------------------------------------
SCAN_COLUMNS = COLUMNS + ["Severity", "Status"]
SCAN_WIDTHS = WIDTHS + [11, 12]
SEV_FILL = {"Critical": "F6D5D2", "Major": "FBE7C6", "Minor": "E4E9EF"}
_TAB = {"Raw Data": "Raw Data", "SDTM": "SDTM", "ADaM": "ADaM", "TLF": "TLF",
        "SAP": "SAP", "CDISC Standards": "CDISC", "Mock Shells": "Mock Shells"}


def status_of(f):
    """Confirmed: code proved it. To check: code noticed it and could not."""
    if f.decision == "accepted":
        return "Kept"
    if f.decision == "rejected":
        return "Rejected"
    if f.confidence == "certain":
        return "Confirmed"
    return "Kept by model" if f.detail.get("judged") else "To check"


def _scan_rows(ws, findings, with_layer=False):
    cols = (["Layer"] if with_layer else []) + SCAN_COLUMNS
    ws.append(cols)
    for f in findings:
        row = [f.area, f.finding, f.output_file, f.what_is_wrong, f.evidence,
               f.severity, status_of(f)]
        ws.append(([f.layer] if with_layer else []) + row)
    widths = ([14] if with_layer else []) + SCAN_WIDTHS
    _style_sheet(ws, len(cols), len(findings), widths)
    sev_col = cols.index("Severity") + 1
    for r in range(2, len(findings) + 2):
        c = ws.cell(row=r, column=sev_col)
        c.fill = PatternFill("solid", fgColor=SEV_FILL.get(c.value, "FFFFFF"))
        c.font = Font(size=10, bold=True)


def _plain_sheet(ws, header, rows, widths):
    ws.append(header)
    for r in rows:
        ws.append(r)
    _style_sheet(ws, len(header), len(rows), widths)


def _compare_sheet(ws, diffs):
    rows = []
    for d in diffs:
        if d.status in ("new", "removed", "same"):
            rows.append([d.new_file or d.old_file, d.old_file, d.status,
                         d.summary(), "", "", ""])
            continue
        for c in d.changes:
            rows.append([d.new_file, d.old_file, "changed", c.kind, c.where,
                         c.old, c.new])
    _plain_sheet(ws, ["Output", "Previous file", "Status", "What changed",
                      "Where", "Previous", "Now"], rows,
                 [40, 34, 10, 14, 34, 40, 40])


_VERDICT = {"confirmed": "Confirmed", "likely_false_alarm": "Likely false alarm",
            "needs_person": "Needs a person"}


def _agent_sheets(wb, findings, agent):
    fixes, queries = [], []
    for f in findings:
        e = agent.get(f.key)
        if not e:
            continue
        for x in e.get("fixes", []):
            fixes.append([f.layer, f.severity, f.finding, f.output_file,
                          _VERDICT.get(e.get("verdict"), ""), x["kind"],
                          x["text"], e.get("by", "")])
        for q in e.get("queries", []):
            queries.append([f.layer, f.severity, f.finding, q["to"],
                            q["subject"], q["text"], e.get("by", "")])
    if fixes:
        _plain_sheet(wb.create_sheet("Agent Fixes"),
                     ["Layer", "Severity", "Finding", "Output / File",
                      "Agent verdict", "Fix type", "Suggested fix", "By"],
                     fixes, [13, 10, 40, 26, 16, 10, 70, 14])
    if queries:
        _plain_sheet(wb.create_sheet("Queries"),
                     ["Layer", "Severity", "Finding", "To", "Subject",
                      "Query", "By"], queries, [13, 10, 40, 16, 40, 80, 14])


def write_scan_workbook(res, path, not_checked=None, agent=None):
    """Summary tab, an All findings tab, then one tab per layer."""
    from .flags import LAYERS
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    counts = res.counts()
    ws.append(["Clinical Review Assistant - scan of %s" % res.study])
    ws.append(["Run on %s. Model: %s." % (res.run_at, res.model_mode)])
    ws.append([])
    ws.append(["Layer", "Critical", "Major", "Minor", "Total", "Inputs read"])
    inputs = {"Raw Data": res.inputs.get("raw", 0), "SDTM": res.inputs.get("sdtm", 0),
              "ADaM": res.inputs.get("adam", 0), "TLF": res.inputs.get("tlf", 0),
              "SAP": res.inputs.get("sap", 0),
              "CDISC Standards": res.inputs.get("sdtm", 0) + res.inputs.get("adam", 0),
              "Mock Shells": res.inputs.get("shells", 0)}
    for layer in LAYERS:
        c = counts.get(layer, {})
        ws.append([layer, c.get("Critical", 0), c.get("Major", 0),
                   c.get("Minor", 0), sum(c.values()), inputs[layer]])
    tot = [sum(counts.get(l, {}).get(s, 0) for l in LAYERS)
           for s in ("Critical", "Major", "Minor")]
    ws.append(["Total"] + tot + [sum(tot), ""])
    head = 4
    thin = Side(style="thin", color=GREY)
    for r in range(head, head + len(LAYERS) + 2):
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.font = Font(size=10, bold=(r in (head, head + len(LAYERS) + 1)
                                            or c == 1),
                             color="FFFFFF" if r == head else "000000")
            if r == head:
                cell.fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(size=14, bold=True, color=NAVY)
    ws["A2"].font = Font(size=10, italic=True)

    r = head + len(LAYERS) + 3
    notes = [("Status", ""),
             ("Confirmed", "Code proved it. Every number in Evidence was counted "
                           "by code, never written by a model."),
             ("To check", "Code noticed it but cannot prove it alone. Look at it, "
                          "or turn the model on to judge it."),
             ("Kept / Rejected", "A person decided, in the app."),
             ("", ""), ("What happened", "")]
    notes += [("", l) for l in res.log]
    if not_checked:
        notes += [("", ""), ("NOT checked", "")] + [("", l) for l in not_checked]
    for a, b in notes:
        ws.cell(row=r, column=1, value=a).font = Font(size=10, bold=True)
        cell = ws.cell(row=r, column=2, value=b)
        cell.font = Font(size=10)
        cell.alignment = Alignment(wrap_text=False, vertical="top")
        r += 1
    for col, w in zip("ABCDEF", (22, 11, 11, 11, 11, 13)):
        ws.column_dimensions[col].width = w

    _scan_rows(wb.create_sheet("All Findings"), res.findings, with_layer=True)
    for layer in LAYERS:
        _scan_rows(wb.create_sheet(_TAB[layer]), res.by_layer(layer))
    if res.tlf_compare:
        _compare_sheet(wb.create_sheet("TLF Compare"), res.tlf_compare)
    if agent:
        _agent_sheets(wb, res.findings, agent)
    wb.save(path)
    return path
