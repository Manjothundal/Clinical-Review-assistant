# -*- coding: utf-8 -*-
"""
The full scan finds every problem planted in the demo study, and puts each one
on the right tab. If a check breaks, the test names it.
"""
import os
import subprocess
import sys
import pytest

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)


@pytest.fixture(scope="module")
def result():
    os.chdir(HERE)
    subprocess.run([sys.executable, "make_sample_study.py"], check=True,
                   capture_output=True)
    from cra.scan_study import scan
    return scan(os.path.join("studies", "DEMO-BE-101"))


def test_every_planted_problem_is_found(result):
    from make_sample_study import PLANTED
    found = {f.found_by for f in result.findings}
    missing = [k for k in PLANTED if k not in found]
    assert not missing, "not found: %s" % missing


@pytest.mark.parametrize("check,layer", [
    ("raw_subject_not_in_sdtm", "Raw Data"), ("raw_vs_sdtm_count", "Raw Data"),
    ("sdtm_subject_not_in_dm", "SDTM"), ("sdtm_study_day", "SDTM"),
    ("adam_chg", "ADaM"), ("adam_saffl_no_dose", "ADaM"),
    ("tlf_n_above_adam", "TLF"), ("sap_population_n", "SAP"),
    ("cdisc_ct", "CDISC Standards"), ("cdisc_iso8601", "CDISC Standards"),
    ("shell_not_delivered", "Mock Shells"),
])
def test_each_lands_on_its_own_tab(result, check, layer):
    layers = {f.layer for f in result.findings if f.found_by == check}
    assert layers == {layer}


def test_every_finding_has_evidence_and_a_layer(result):
    from cra.flags import LAYERS
    for f in result.findings:
        assert f.evidence.strip()
        assert f.layer in LAYERS


def test_numbers_in_the_headline_are_in_the_evidence(result):
    """The rule the model is held to holds for the code's own words too."""
    from cra.review import numbers_not_in_evidence
    for f in result.findings:
        if f.found_by.startswith(("raw_", "sdtm_", "adam_", "sap_", "cdisc_",
                                  "shell_", "tlf_", "spec_")):
            assert not numbers_not_in_evidence(f.finding, f.evidence), f.finding


def test_rejected_finding_stays_rejected(result, tmp_path):
    from cra.scan_study import scan, save_decision, decisions_path
    folder = result.folder
    first = result.findings[0]
    try:
        save_decision(folder, first.key, "rejected")
        again = scan(folder)
        assert first.key not in {f.key for f in again.findings}
    finally:
        os.remove(decisions_path(folder))


def test_sheet_has_one_tab_per_layer(result, tmp_path):
    import openpyxl
    from cra.excel_out import write_scan_workbook
    p = write_scan_workbook(result, str(tmp_path / "scan.xlsx"))
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Summary", "All Findings", "Raw Data", "SDTM",
                             "ADaM", "TLF", "SAP", "CDISC", "Mock Shells",
                             "TLF Compare"]
    assert wb["All Findings"].max_row == len(result.findings) + 1
    head = [c.value for c in wb["SDTM"][1]]
    assert head[:5] == ["Area", "Finding", "Output / File", "What is wrong",
                        "Evidence"]


def test_upload_sorting():
    from cra.study import classify
    assert classify("adsl.xpt") == "adam"
    assert classify("dm.xpt") == "sdtm"
    assert classify("suppae.xpt") == "sdtm"
    assert classify("ae_raw.csv") == "raw"
    assert classify("EDC_DEMOG.sas7bdat") == "raw"
    assert classify("t_ae.rtf") == "tlf"
    assert classify("Mock_Shells_v2.docx") == "shells"
    assert classify("ABC_SAP_v1.pdf") == "sap"
    assert classify("spec.txt") == "sap"
    assert classify("ADaM_specs.xlsx") == "specs"
    assert classify("delivery/sdtm/xx.xpt") == "sdtm"


def test_only_candidates_go_to_the_model(result, monkeypatch):
    """Certain findings are not sent: that is where the tokens are saved."""
    import json
    from cra import review
    sent = {}

    def fake(system, user, mode=None):
        sent["items"] = json.loads(user.split("ITEMS_JSON:")[1])
        return "[]"
    monkeypatch.setattr(review, "ask_model", fake)
    from cra.scan_study import scan
    r = scan(result.folder, use_model=True, mode="openai")
    n_cand = sum(1 for f in result.findings if f.confidence == "candidate")
    assert len(sent["items"]) == n_cand
    assert all(i["conf"] == "candidate" for i in sent["items"])
    # model said nothing about them, so they stay, marked To check
    assert len(r.findings) == len(result.findings)


def test_model_failure_loses_nothing(result, monkeypatch):
    from cra import review

    def boom(*a, **k):
        raise RuntimeError("no network")
    monkeypatch.setattr(review, "ask_model", boom)
    from cra.scan_study import scan
    r = scan(result.folder, use_model=True, mode="openai")
    assert len(r.findings) == len(result.findings)


def test_old_vs_new_tlf(result):
    """Planted: 1 cell, 1 title, 1 footnote, 1 program change, 1 new, 1 gone.
    A new run date on its own is not a change."""
    from cra.compare_tlf import totals
    t = totals(result.tlf_compare)
    assert (t["changed"], t["new"], t["removed"], t["cells"]) == (4, 1, 1, 1)
    by = {d.new_file or d.old_file: d for d in result.tlf_compare}
    c = by["Table_12_5_Urinalysis.rtf"].changes[0]
    assert (c.kind, c.old, c.new) == ("cell", "48", "12")
    assert by["Table_12_7_Human_Factors.rtf"].changes[0].kind == "title"
    assert by["Appendix_16_2_5_3_Vital_Signs.rtf"].changes[0].kind == "source"
    assert by["Table_12_9_Chemistry.rtf"].status == "removed"
    assert by["SbjGraphs_IBA_adj_6050.rtf"].status == "new"
    assert by["PK_DesStats_EPIN_adj_6050.rtf"].status == "same"
